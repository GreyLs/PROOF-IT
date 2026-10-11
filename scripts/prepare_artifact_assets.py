from pathlib import Path
import urllib.request,gzip,shutil
from fontTools.ttLib import TTFont
from fontTools.varLib.instancer import instantiateVariableFont
from pptx import Presentation
ROOT=Path(__file__).resolve().parents[1]
out=ROOT/'research/fonts';out.mkdir(exist_ok=True)
p=out/'Montserrat-VF.ttf'
if not p.exists():p.write_bytes(urllib.request.urlopen('https://raw.githubusercontent.com/google/fonts/main/ofl/montserrat/Montserrat%5Bwght%5D.ttf',timeout=45).read())
for weight,name in [(400,'Regular'),(700,'Bold')]:
    f=TTFont(p);f=instantiateVariableFont(f,{'wght':weight},inplace=True)
    for platform,encoding,language in [(3,1,0x409),(1,0,0)]:
        f['name'].setName(f'MontserratStatic-{name}',6,platform,encoding,language)
        f['name'].setName(f'Montserrat Static {name}',4,platform,encoding,language)
    f.save(out/f'Montserrat-{name}.ttf')
(out/'OFL.txt').write_bytes(urllib.request.urlopen('https://raw.githubusercontent.com/google/fonts/main/ofl/montserrat/OFL.txt',timeout=45).read())
(ROOT/'web/public/brand/OFL-JetBrainsMono.txt').write_bytes(urllib.request.urlopen('https://raw.githubusercontent.com/JetBrains/JetBrainsMono/master/OFL.txt',timeout=45).read())
with (ROOT/'data/relevance_labels.csv').open('rb') as src,gzip.open(ROOT/'data/relevance_labels.csv.gz','wb',compresslevel=9) as dest:shutil.copyfileobj(src,dest)
prs=Presentation(next((ROOT/'research/case').rglob('*.pptx')))
for shape in prs.slides[6].shapes:
    if shape.shape_type==13:
        (ROOT/'research/template-logo.png').write_bytes(shape.image.blob);break
print('Fonts, template logo and compressed oracle ready')
